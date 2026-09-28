import xlsxwriter 
from openpyxl import load_workbook 
import os

class nhanvien:
    def __init__(self,manv,ten,tuoi):
        self.manv=manv
        self.ten=ten
        self.tuoi=tuoi
    def Xuat(self):
        print("Mã: ",self.manv)
        print("Tên: ",self.ten)
        print("Tuổi: ",self.tuoi)

def LuuFile(duongdan,DSnhanvien):
    # Tạo một file excel cùng 1 sheet 
    workbook = xlsxwriter.Workbook(duongdan) 
    worksheet = workbook.add_worksheet() 
    
    # thiết lập các cột cho file 
    worksheet.set_column('A:A', 15) 
    worksheet.set_column('B:B', 15) 
    worksheet.set_column('C:C', 15) 
    worksheet.set_column('D:D', 15) 
    
    # định dạng tiêu đề cột in đậm 
    bold = workbook.add_format({'bold': True}) 
    
    # thêm dòng tiêu đề và định dạng in đậm
    worksheet.write('A1', 'STT',bold) 
    worksheet.write('B1', 'MÃ',bold) 
    worksheet.write('C1', 'TÊN',bold) 
    worksheet.write('D1', 'TUỔI',bold)
    
    # thêm dữ liệu
    dong = 2
    stt = 1

    for nv in DSnhanvien:
        worksheet.write(dong - 1, 0, stt)
        worksheet.write(dong - 1, 1, nv.manv)
        worksheet.write(dong - 1, 2, nv.ten)
        worksheet.write(dong - 1, 3, nv.tuoi)

        dong += 1
        stt += 1

    workbook.close()

def DocFile(duongdan):
    wb = load_workbook(duongdan)
    print (wb.sheetnames) 
    ws = wb[wb.sheetnames[0]] 
    DSnhanvien = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        stt, manv, ten, tuoi = row

        nv = nhanvien(manv, ten, tuoi)
        DSnhanvien.append(nv)

    return DSnhanvien

def main():
    duongdan=os.path.join(os.path.dirname(__file__), "nhanvien.xlsx")

    if os.path.exists(duongdan):
        DSnhanvien = DocFile(duongdan)
    else:
        DSnhanvien = []

    while(True):
        print("0. Thoát")
        print("1. Thêm")
        print("2. Xuất danh sách")
        print("3. Sắp xếp")
        print("4. Lưu")
        print("Mục bạn chọn: ",end='')
        muc=int(input())

        if muc==0:
            break
        elif muc==1:
            manv=input("Mã nhân viên: ")
            ten=input("Tên: ")
            tuoi=input("Tuổi: ")
            nv=nhanvien(manv,ten,tuoi)
            DSnhanvien.append(nv)
        elif muc==2:
            stt=1
            for nv in DSnhanvien:
                print("Nhân viên thứ ",stt)
                nv.Xuat()
                stt+=1
        elif muc==3:
            for i in range(len(DSnhanvien)-1):
                for j in range(i+1,len(DSnhanvien)):
                    if float(DSnhanvien[i].tuoi) > float(DSnhanvien[j].tuoi):
                        DSnhanvien[i], DSnhanvien[j]=DSnhanvien[j], DSnhanvien[i]
        elif muc==4:
            LuuFile(duongdan, DSnhanvien)
            print("Đã lưu file!")

if __name__=="__main__":
    main()