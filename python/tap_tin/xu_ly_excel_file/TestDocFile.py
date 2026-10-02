# Sử dụng thư viện openpyxl để đọc file excel đã được lưu từ file TestLuuFile.py

from openpyxl import load_workbook 
import os

duongdan=os.path.join(os.path.dirname(__file__), "demo.xlsx")
wb = load_workbook(duongdan)
print (wb.sheetnames) 
ws = wb[wb.sheetnames[0]] 
for row in ws.values: 
   for value in row: 
     print(value,"\t",end='') 
   print("") 
